# -*- coding: utf-8 -*-
"""
sync_rem.py — Relevamiento de Expectativas de Mercado (REM) del BCRA para el Ministro y Argento.

Busca el Excel del último REM en la web del BCRA, lo guarda en macro/rem/ y deja macro/rem.json con:
  meta: url, fecha de descarga, hojas, diagnóstico de la búsqueda
  muestra: primeras filas de cada hoja (para ajustar el lector)
  ipc / tamar / tc / pib: medianas por período cuando el lector las reconoce
Corre en GitHub Actions (tiene red). Sin dependencias fuera de openpyxl.
"""
from __future__ import annotations
import os, re, json, io, urllib.request, urllib.parse
from datetime import date, datetime, timezone

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "macro")
HDR = {"User-Agent": "Mozilla/5.0 (rem-sync)"}
PAGINAS = [
    "https://www.bcra.gob.ar/PublicacionesEstadisticas/Relevamiento_Expectativas_de_Mercado.asp",
    "https://www.bcra.gob.ar/relevamiento-de-expectativas-de-mercado-rem/",
    "https://www.bcra.gob.ar/publicaciones-estadisticas/relevamiento-de-expectativas-de-mercado-rem/",
]
MESES = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]

def get(url, timeout=60):
    req = urllib.request.Request(url, headers=HDR)
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return r.read(), r.geturl()

def buscar_links(diag):
    links = []
    for p in PAGINAS:
        try:
            html, final = get(p)
            t = html.decode("utf-8", "ignore")
            hs = re.findall(r'href=["\']([^"\']+\.xlsx?)["\']', t, re.I)
            hs = [urllib.parse.urljoin(final, h) for h in hs]
            rem = [h for h in hs if re.search(r"rem|expectativ", h, re.I)]
            diag.append({"pagina": p, "ok": True, "xlsx": len(hs), "rem": rem[:10]})
            links += rem
        except Exception as e:
            diag.append({"pagina": p, "ok": False, "error": repr(e)[:200]})
    # patrones conocidos por si la página cambió: últimos 4 meses
    hoy = date.today()
    for k in range(4):
        y, m = hoy.year, hoy.month - k
        while m <= 0: m += 12; y -= 1
        aamm = f"{y % 100:02d}{m:02d}"
        for pat in (f"https://www.bcra.gob.ar/Pdfs/PublicacionesEstadisticas/REM{aamm}%20Tablas%20web.xlsx",
                    f"https://www.bcra.gob.ar/Pdfs/PublicacionesEstadisticas/REM{aamm}_Tablas%20web.xlsx",
                    f"https://www.bcra.gob.ar/Pdfs/PublicacionesEstadisticas/tablas-relevamiento-expectativas-mercado-{MESES[m-1]}-{y}.xlsx"):
            links.append(pat)
    vistos, out = set(), []
    for l in links:
        if l not in vistos: vistos.add(l); out.append(l)
    return out

def muestra(wb, filas=45, cols=16):
    m = {}
    for ws in wb.worksheets:
        rows = []
        for r in ws.iter_rows(min_row=1, max_row=filas, max_col=cols, values_only=True):
            if any(v not in (None, "") for v in r):
                rows.append([("" if v is None else (v.isoformat()[:10] if hasattr(v, "isoformat") else str(v)[:60])) for v in r])
        m[ws.title] = rows
    return m

def leer_medianas(wb):
    """Lector genérico: busca filas con el nombre de la variable y toma (período, mediana) de las filas siguientes."""
    claves = {"ipc": r"precios minoristas|ipc", "tamar": r"tamar", "badlar": r"badlar", "tc": r"tipo de cambio",
              "pib": r"pib", "desempleo": r"desocupaci", "resultado_fiscal": r"resultado primario|fiscal"}
    out = {}
    for ws in wb.worksheets:
        filas = list(ws.iter_rows(values_only=True))
        for i, r in enumerate(filas):
            celda = " ".join(str(v) for v in r[:3] if v is not None).lower()
            for k, pat in claves.items():
                if k in out or not re.search(pat, celda): continue
                # encabezado con "mediana" en las filas cercanas
                for j in range(i, min(i + 4, len(filas))):
                    hdr = [str(v).lower() if v is not None else "" for v in filas[j]]
                    if any("mediana" in h for h in hdr):
                        cm = next(c for c, h in enumerate(hdr) if "mediana" in h)
                        serie = []
                        for r2 in filas[j + 1:j + 30]:
                            per = r2[0] if r2 and r2[0] not in (None, "") else (r2[1] if len(r2) > 1 else None)
                            val = r2[cm] if cm < len(r2) else None
                            if per in (None, "") and val in (None, ""): break
                            if isinstance(val, (int, float)):
                                serie.append([per.isoformat()[:10] if hasattr(per, "isoformat") else str(per), float(val)])
                        if serie: out[k] = {"hoja": ws.title, "fila": i + 1, "serie": serie}
                        break
    return out

def main():
    import openpyxl
    os.makedirs(os.path.join(OUT, "rem"), exist_ok=True)
    diag, res = [], {"meta": {"updatedAt": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%MZ"), "by": "sync_rem.py (GitHub Actions)"}}
    url_ok, data = None, None
    for l in buscar_links(diag):
        try:
            b, final = get(l)
            if b[:2] == b"PK":   # xlsx
                url_ok, data = final, b; break
            diag.append({"link": l, "no_xlsx": len(b)})
        except Exception as e:
            diag.append({"link": l, "error": repr(e)[:160]})
    res["meta"]["diag"] = diag
    if data:
        open(os.path.join(OUT, "rem", "rem_ultimo.xlsx"), "wb").write(data)
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        res["meta"].update({"url": url_ok, "hojas": wb.sheetnames, "bytes": len(data)})
        res["muestra"] = muestra(wb)
        res["medianas"] = leer_medianas(wb)
    else:
        res["meta"]["error"] = "no se encontró el Excel del REM"
    json.dump(res, open(os.path.join(OUT, "rem.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("rem.json:", res["meta"].get("url"), "| medianas:", list((res.get("medianas") or {}).keys()))

if __name__ == "__main__":
    main()
